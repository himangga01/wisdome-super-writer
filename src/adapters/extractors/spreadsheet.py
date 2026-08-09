from __future__ import annotations

import csv
from pathlib import Path
from typing import Any, Mapping

from adapters.extractors.base import ExtractorError, GenericEvidenceRecord, GenericExtractionOutput


class SpreadsheetExtractor:
    engine = "spreadsheet_parser"

    def __init__(self, config: Mapping[str, Any] | None = None) -> None:
        self.config = dict(config or {})
        self.max_file_bytes = int(self.config.get("max_file_bytes", 100 * 1024 * 1024))
        self.max_sheets = int(self.config.get("max_sheets", 100))
        self.max_cells = int(self.config.get("max_cells", 1_000_000))

    def extract(self, path: Path) -> GenericExtractionOutput:
        if path.stat().st_size > self.max_file_bytes:
            raise ExtractorError("spreadsheet_limit_exceeded", "Spreadsheet exceeds the byte limit")
        if path.suffix.lower() in {".csv", ".tsv"}:
            return self._extract_delimited(path)
        if path.suffix.lower() not in {".xlsx", ".xlsm"}:
            raise ExtractorError("spreadsheet_format_unsupported", "Only XLSX/XLSM/CSV/TSV are supported")
        try:
            import openpyxl
        except ImportError as exc:
            raise ExtractorError("spreadsheet_dependency_missing", "openpyxl is not installed") from exc
        try:
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
        except Exception as exc:
            raise ExtractorError("spreadsheet_corrupt", "Spreadsheet parser rejected the input") from exc
        records: list[GenericEvidenceRecord] = []
        cells_seen = 0
        try:
            if len(workbook.sheetnames) > self.max_sheets:
                raise ExtractorError("spreadsheet_limit_exceeded", "Spreadsheet has too many sheets")
            for worksheet in workbook.worksheets:
                for row in worksheet.iter_rows():
                    values = []
                    nonempty = []
                    for cell in row:
                        cells_seen += 1
                        if cells_seen > self.max_cells:
                            raise ExtractorError("spreadsheet_limit_exceeded", "Spreadsheet has too many cells")
                        value = cell.value
                        values.append(value)
                        if value is not None:
                            nonempty.append((cell.coordinate, value, cell.data_type))
                    if not nonempty:
                        continue
                    first, last = nonempty[0][0], nonempty[-1][0]
                    records.append(GenericEvidenceRecord(
                        kind="spreadsheet",
                        locator_type="spreadsheet_cell",
                        locator={
                            "locator_type": "spreadsheet_cell",
                            "sheet_name": worksheet.title,
                            "cell_range": f"{first}:{last}",
                            "table_name": None,
                        },
                        text=" | ".join("" if value is None else str(value) for value in values),
                        structured_data={
                            "cells": [
                                {"coordinate": coordinate, "value": value, "data_type": data_type}
                                for coordinate, value, data_type in nonempty
                            ]
                        },
                    ))
        finally:
            workbook.close()
        return GenericExtractionOutput(
            engine=self.engine,
            extractor_version="1.0.0",
            validation_mode="deterministic",
            records=records,
            metadata={"sheet_count": len(workbook.sheetnames), "cells_scanned": cells_seen},
        )

    def _extract_delimited(self, path: Path) -> GenericExtractionOutput:
        delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
        records: list[GenericEvidenceRecord] = []
        with path.open("r", encoding="utf-8-sig", newline="") as source:
            reader = csv.reader(source, delimiter=delimiter)
            for row_index, row in enumerate(reader, start=1):
                if row_index * max(1, len(row)) > self.max_cells:
                    raise ExtractorError("spreadsheet_limit_exceeded", "Delimited data has too many cells")
                records.append(GenericEvidenceRecord(
                    kind="spreadsheet",
                    locator_type="spreadsheet_cell",
                    locator={
                        "locator_type": "spreadsheet_cell",
                        "sheet_name": "Sheet1",
                        "cell_range": f"A{row_index}:{self._column_name(max(1, len(row)))}{row_index}",
                        "table_name": None,
                    },
                    text=" | ".join(row),
                    structured_data={"row_index": row_index, "values": row},
                ))
        return GenericExtractionOutput(
            engine=self.engine,
            extractor_version="1.0.0",
            validation_mode="deterministic",
            records=records,
            metadata={"row_count": len(records), "delimiter": delimiter},
        )

    @staticmethod
    def _column_name(index: int) -> str:
        result = ""
        while index:
            index, remainder = divmod(index - 1, 26)
            result = chr(65 + remainder) + result
        return result

